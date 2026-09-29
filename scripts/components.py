"""research-slides components: every drawn shape goes through here.

Colour rule (lab rule): shapes of the same kind always share one style.
A different colour appears only when the shape has a different *role*
(the paper's key element, a problem/emphasis, a verified result, a note).
So callers choose a `kind`, never a colour.
"""
import re
import subprocess
import tempfile
from pathlib import Path

from lxml import etree
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_MARKER_STYLE
from pptx.enum.dml import MSO_THEME_COLOR
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

THEME = {
    "tx1": MSO_THEME_COLOR.TEXT_1, "bg1": MSO_THEME_COLOR.BACKGROUND_1,
    "tx2": MSO_THEME_COLOR.TEXT_2, "bg2": MSO_THEME_COLOR.BACKGROUND_2,
    "accent1": MSO_THEME_COLOR.ACCENT_1, "accent2": MSO_THEME_COLOR.ACCENT_2,
    "accent3": MSO_THEME_COLOR.ACCENT_3, "accent4": MSO_THEME_COLOR.ACCENT_4,
    "accent5": MSO_THEME_COLOR.ACCENT_5, "accent6": MSO_THEME_COLOR.ACCENT_6,
}

# --------------------------------------------------------------------------
# Style tables (the only place colours are chosen)
# --------------------------------------------------------------------------
BOX = {
    # kind:     fill,          line,          line_pt, text colour, bold
    "default": ("accent1:12", None,          0,   "accent1", True),   # ordinary module / step
    "key":     ("accent1",    None,          0,   "bg1",     True),   # the paper's proposed element
    "emph":    ("bg1",        "accent2",     2.5, "accent2", True),   # problem, pitfall, key claim
    "ok":      ("accent4:35", None,          0,   "accent1", True),   # verified / correct / pass
    "note":    ("accent5:30", None,          0,   "tx1:75",  False),  # low priority, reference
    "muted":   ("bg2",        None,          0,   "tx1:75",  False),  # baseline, inactive
    "outline": ("bg1",        "accent1:40",  1.25, "tx1",    False),  # container / group frame
}
ARROW = {"default": ("accent6", 3.0), "key": ("accent1", 3.5), "emph": ("accent2", 3.5)}
HIGHLIGHT = ("accent2", 3.0)            # red frame over a figure / equation region
SERIES = {"ours": "accent2", "baseline": "accent6", "alt": "accent3", "alt2": "accent1",
          "ok": "accent4", "note": "accent5"}
TEXT = {"body": ("tx1", False), "muted": ("tx1:75", False), "key": ("accent2", True),
        "struct": ("accent1", True), "caption": ("tx1:75", False)}


def set_color(cf, spec):
    """Apply a colour spec ('accent1', 'accent1:12', 'tx1:75', '#RRGGBB') to a ColorFormat."""
    if spec.startswith("#"):
        cf.rgb = RGBColor.from_string(spec[1:])
        return
    name, _, pct = spec.partition(":")
    cf.theme_color = THEME[name]
    if pct:
        cf.brightness = round((100 - int(pct)) / 100, 3)


def _fill(shape, spec):
    if spec is None:
        shape.fill.background()
    else:
        shape.fill.solid()
        set_color(shape.fill.fore_color, spec)


def _line(shape, spec, pt):
    if spec is None:
        shape.line.fill.background()
    else:
        shape.line.fill.solid()
        set_color(shape.line.color, spec)
        shape.line.width = Pt(pt)


def _no_shadow(shape):
    sppr = shape._element.spPr
    if sppr.find(qn("a:effectLst")) is None:
        etree.SubElement(sppr, qn("a:effectLst"))
    style = shape._element.find(qn("p:style"))
    if style is not None:  # theme effect style would add a drop shadow
        eff = style.find(qn("a:effectRef"))
        if eff is not None:
            eff.set("idx", "0")


# --------------------------------------------------------------------------
# Inline markup:  **key**  -> Dusty Cedar bold   (핵심어)
#                 __term__ -> Navy Peony bold     (구조 용어)
# --------------------------------------------------------------------------
TOKEN = re.compile(r"(\*\*.+?\*\*|__.+?__)")


def add_runs(paragraph, s, size=None, color=None, bold=None):
    for part in TOKEN.split(s):
        if not part:
            continue
        r = paragraph.add_run()
        if part.startswith("**") and part.endswith("**"):
            r.text = part[2:-2]
            set_color(r.font.color, "accent2")
            r.font.bold = True
        elif part.startswith("__") and part.endswith("__"):
            r.text = part[2:-2]
            set_color(r.font.color, "accent1")
            r.font.bold = True
        else:
            r.text = part
            if color:
                set_color(r.font.color, color)
            if bold is not None:
                r.font.bold = bold
        if size:
            r.font.size = Pt(size)


def split_level(item):
    """'>> text' -> (2, 'text');  '> text' -> (1, ..);  'text' -> (0, ..)."""
    m = re.match(r"^(>{1,2})\s*", item)
    if not m:
        return 0, item
    return len(m.group(1)), item[m.end():]


BULLET = {0: ("Wingdings", "\uf0fc", 0.5), 1: ("Wingdings", "\uf0d8", 1.0), 2: ("Arial", "\u2013", 1.45)}


def _bullet(p, lvl, size):
    """Explicit lab bullets (check / arrow / dash) for free text boxes."""
    font, char, marl = BULLET[min(lvl, 2)]
    pPr = p._p.get_or_add_pPr()
    ind = 0.42 if lvl == 0 else 0.4
    pPr.set("marL", str(int(Inches(marl))))
    pPr.set("indent", str(-int(Inches(ind))))
    for tag in ("a:buClr", "a:buSzPct", "a:buFont", "a:buChar", "a:buNone"):
        for e in pPr.findall(qn(tag)):
            pPr.remove(e)
    buClr = etree.SubElement(pPr, qn("a:buClr"))
    etree.SubElement(buClr, qn("a:schemeClr"), val="accent1")
    etree.SubElement(pPr, qn("a:buSzPct"), val="90000")
    etree.SubElement(pPr, qn("a:buFont"), typeface=font, pitchFamily="2", charset="2")
    etree.SubElement(pPr, qn("a:buChar"), char=char)


def fill_text(tf, items, size=None, color=None, bold=None, levels=True, align=None,
              bullets=False):
    """Write a string or list of strings into a text frame (placeholder or box).

    bullets=True draws the lab bullets explicitly (for free text boxes; layout
    placeholders already carry them in their list styles)."""
    if isinstance(items, str):
        items = items.split("\n")
    tf.clear()
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        lvl, s = split_level(item) if levels else (0, item)
        p.level = lvl
        if bullets:
            _bullet(p, lvl, size)
        if align:
            p.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[align]
        add_runs(p, s, size=size, color=color, bold=bold)


# --------------------------------------------------------------------------
# Shapes
# --------------------------------------------------------------------------
def box(slide, x, y, w, h, text="", kind="default", size=24, align="c", shape="round"):
    fill, line, lpt, tcolor, bold = BOX[kind]
    geom = {"round": MSO_SHAPE.ROUNDED_RECTANGLE, "rect": MSO_SHAPE.RECTANGLE,
            "pill": MSO_SHAPE.ROUNDED_RECTANGLE, "oval": MSO_SHAPE.OVAL}[shape]
    sp = slide.shapes.add_shape(geom, Inches(x), Inches(y), Inches(w), Inches(h))
    if shape == "round":
        sp.adjustments[0] = min(0.5, 0.14 / max(min(w, h), 0.01))  # ~0.14 in corner radius
    elif shape == "pill":
        sp.adjustments[0] = 0.5
    _fill(sp, fill)
    _line(sp, line, lpt)
    _no_shadow(sp)
    tf = sp.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    for side in ("margin_left", "margin_right"):
        setattr(tf, side, Inches(0.15))
    fill_text(tf, text, size=size, color=tcolor, bold=bold, levels=False, align=align)
    return sp


def arrow(slide, x1, y1, x2, y2, kind="default", head="end"):
    color, pt = ARROW[kind]
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1),
                                   Inches(x2), Inches(y2))
    c.line.fill.solid()
    set_color(c.line.color, color)
    c.line.width = Pt(pt)
    ln = c.line._get_or_add_ln()
    if head in ("end", "both"):
        etree.SubElement(ln, qn("a:tailEnd"), type="triangle", w="med", len="lg")
    if head in ("start", "both"):
        etree.SubElement(ln, qn("a:headEnd"), type="triangle", w="med", len="lg")
    return c


def highlight(slide, x, y, w, h):
    color, pt = HIGHLIGHT
    sp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sp.adjustments[0] = 0.08
    _fill(sp, None)
    _line(sp, color, pt)
    _no_shadow(sp)
    return sp


def textbox(slide, x, y, w, h, text, style="body", size=24, align="l", anchor="t", levels=True,
            bullets=False):
    color, bold = TEXT[style]
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[anchor]
    fill_text(tf, text, size=size, color=color, bold=bold, levels=levels, align=align,
              bullets=bullets)
    if bullets:  # second-level lines in muted grey, like the placeholders
        for p in tf.paragraphs:
            if p.level:
                for r in p.runs:
                    if r.font.color and r.font.color.type is not None and \
                            r.font.color.theme_color == THEME["tx1"]:
                        set_color(r.font.color, "tx1:75")
    return tb


def flow(slide, x, y, w, h, items, key=(), emph=(), gap=0.8, size=24, vertical=False):
    """Boxes joined by arrows. key/emph: indices drawn with that kind."""
    n = len(items)
    if vertical:
        bh = (h - gap * (n - 1)) / n
        for i, it in enumerate(items):
            by = y + i * (bh + gap)
            k = "key" if i in key else "emph" if i in emph else "default"
            box(slide, x, by, w, bh, it, kind=k, size=size)
            if i < n - 1:
                arrow(slide, x + w / 2, by + bh + 0.06, x + w / 2, by + bh + gap - 0.06)
    else:
        bw = (w - gap * (n - 1)) / n
        for i, it in enumerate(items):
            bx = x + i * (bw + gap)
            k = "key" if i in key else "emph" if i in emph else "default"
            box(slide, bx, y, bw, h, it, kind=k, size=size)
            if i < n - 1:
                arrow(slide, bx + bw + 0.06, y + h / 2, bx + bw + gap - 0.06, y + h / 2)


def chevrons(slide, x, y, w, h, items, key=(), size=22):
    n = len(items)
    overlap = h * 0.25
    cw = (w + overlap * (n - 1)) / n
    for i, it in enumerate(items):
        sp = slide.shapes.add_shape(MSO_SHAPE.CHEVRON if i else MSO_SHAPE.PENTAGON,
                                    Inches(x + i * (cw - overlap)), Inches(y), Inches(cw), Inches(h))
        kind = "key" if i in key else "default"
        fill, _, _, tcolor, bold = BOX[kind]
        _fill(sp, fill)
        _line(sp, "bg1", 2.0)
        _no_shadow(sp)
        sp.text_frame.word_wrap = True
        fill_text(sp.text_frame, it, size=size, color=tcolor, bold=True, levels=False, align="c")


def table(slide, x, y, w, rows, col_widths=None, size=20, row_h=0.6,
          ours_rows=(), key_cells=(), first_col=True):
    nr, nc = len(rows), len(rows[0])
    gf = slide.shapes.add_table(nr, nc, Inches(x), Inches(y), Inches(w), Inches(row_h * nr))
    tbl = gf.table
    # "No Style, No Grid" so every colour below is ours
    gf._element.graphic.graphicData.tbl.tblPr.find(qn("a:tableStyleId")).text = \
        "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"
    tbl.first_row = True
    tbl.horz_banding = False
    widths = col_widths or [w / nc] * nc
    for j, cw in enumerate(widths):
        tbl.columns[j].width = Inches(cw)
    for i in range(nr):
        tbl.rows[i].height = Inches(row_h)
        for j in range(nc):
            cell = tbl.cell(i, j)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.margin_left = cell.margin_right = Inches(0.12)
            header, ours = i == 0, i in ours_rows
            key = [i, j] in [list(k) for k in key_cells]
            if header:
                color, bold, fill = "bg1", True, "accent1"
            elif j == 0 and first_col:
                color, bold, fill = "tx1", True, "bg2"
            else:
                color, bold, fill = "tx1", False, None
            if (ours or key) and not header:
                color, bold = "accent2", True
            if fill:
                cell.fill.solid()
                set_color(cell.fill.fore_color, fill)
            else:
                cell.fill.background()
            fill_text(cell.text_frame, str(rows[i][j]), size=size, color=color, bold=bold,
                      levels=False, align="l" if j == 0 else "c")
            tcPr = cell._tc.get_or_add_tcPr()
            lnB = etree.Element(qn("a:lnB"), w="12700")
            tcPr.insert(0, lnB)  # borders must precede the cell fill in tcPr
            sf = etree.SubElement(lnB, qn("a:solidFill"))
            sc = etree.SubElement(sf, qn("a:schemeClr"), val="accent1")
            etree.SubElement(sc, qn("a:lumMod"), val="30000")
            etree.SubElement(sc, qn("a:lumOff"), val="70000")
    return gf


def chart(slide, x, y, w, h, categories, series, kind="bar", ylabel=None, legend=True,
          labels=False, size=18, number_format=None):
    """Native (editable) chart. series: [{"name", "values", "role": ours|baseline|alt|alt2}]."""
    cd = CategoryChartData()
    cd.categories = categories
    for s in series:
        cd.add_series(s["name"], s["values"])
    ct = {"bar": XL_CHART_TYPE.COLUMN_CLUSTERED, "hbar": XL_CHART_TYPE.BAR_CLUSTERED,
          "line": XL_CHART_TYPE.LINE_MARKERS}[kind]
    gf = slide.shapes.add_chart(ct, Inches(x), Inches(y), Inches(w), Inches(h), cd)
    ch = gf.chart
    ch.font.size = Pt(size)
    set_color(ch.font.color, "tx1:75")
    ch.has_legend = legend and len(series) > 1
    if ch.has_legend:
        ch.legend.position = XL_LEGEND_POSITION.TOP
        ch.legend.include_in_layout = False
    va = ch.value_axis
    va.has_major_gridlines = True
    va.major_gridlines.format.line.width = Pt(0.75)
    set_color(va.major_gridlines.format.line.color, "tx1:15")
    va.format.line.fill.background()
    if number_format:
        va.tick_labels.number_format = number_format
        va.tick_labels.number_format_is_linked = False
    if ylabel:
        va.has_title = True
        va.axis_title.text_frame.text = ylabel
        va.axis_title.text_frame.paragraphs[0].runs[0].font.size = Pt(size)
    ca = ch.category_axis
    ca.format.line.width = Pt(1)
    set_color(ca.format.line.color, "tx1:40")
    for s_obj, s in zip(ch.plots[0].series, series):
        col = SERIES[s.get("role", "alt")]
        if kind == "line":
            s_obj.format.line.width = Pt(3)
            set_color(s_obj.format.line.color, col)
            s_obj.marker.style = XL_MARKER_STYLE.CIRCLE
            s_obj.marker.size = 9
            s_obj.marker.format.fill.solid()
            set_color(s_obj.marker.format.fill.fore_color, col)
            s_obj.marker.format.line.fill.background()
            s_obj.smooth = False
        else:
            s_obj.format.fill.solid()
            set_color(s_obj.format.fill.fore_color, col)
            s_obj.format.line.fill.background()
    if kind != "line":
        ch.plots[0].gap_width = 60
        ch.plots[0].overlap = -10
    if labels:
        pl = ch.plots[0]
        pl.has_data_labels = True
        pl.data_labels.font.size = Pt(size - 2)
        if number_format:
            pl.data_labels.number_format = number_format
            pl.data_labels.number_format_is_linked = False
    return gf


def image_fit(slide, path, x, y, w, h, crop=None, align="c"):
    """Place an image inside a box without distortion or cropping content."""
    from PIL import Image
    with Image.open(path) as im:
        iw, ih = im.size
    l, t, r, b = crop or (0, 0, 0, 0)
    iw2, ih2 = iw * (1 - l - r), ih * (1 - t - b)
    scale = min(w / iw2, h / ih2)
    dw, dh = iw2 * scale, ih2 * scale
    dx = x + {"l": 0, "c": (w - dw) / 2, "r": w - dw}[align]
    dy = y + (h - dh) / 2
    pic = slide.shapes.add_picture(str(path), Inches(dx), Inches(dy), Inches(dw), Inches(dh))
    if crop:
        pic.crop_left, pic.crop_top, pic.crop_right, pic.crop_bottom = l, t, r, b
    return pic, (dx, dy, dw, dh)


def math_image(tex, size=32, color="#262626", dpi=300):
    """Render a LaTeX-ish formula with matplotlib mathtext; return PNG path."""
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    fig = plt.figure(figsize=(0.01, 0.01))
    fig.text(0, 0, f"${tex}$", fontsize=size, color=color)
    out = Path(tempfile.mkstemp(suffix=".png")[1])
    fig.savefig(out, dpi=dpi, transparent=True, bbox_inches="tight", pad_inches=0.05)
    plt.close(fig)
    return out


def swatch(slide, x, y, w, h, color, label="", sub=""):
    """Palette swatch (guide slides only; content slides use kinds, not colours)."""
    sp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(w))
    sp.adjustments[0] = 0.12
    _fill(sp, color)
    _line(sp, None, 0)
    _no_shadow(sp)
    textbox(slide, x - 0.1, y + w + 0.1, w + 1.2, h - w - 0.1, [label, sub], style="body", size=16,
            levels=False)
