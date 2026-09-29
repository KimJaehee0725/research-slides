#!/usr/bin/env python3
"""Build the research-slides PowerPoint template from scratch.

    python3 scripts/make_template.py [--out assets/research-slides-template.pptx]

Everything that should look the same on every slide lives in the slide
master or a named layout. Content areas are named placeholders, so a deck
builder (human or Claude) picks a layout and fills roles such as
`title`, `figure`, `band` instead of copying and moving shapes.

Re-run this script after editing rs_style.py or the layout table below.
By default the hidden guide slides and one example slide per layout
(examples/template_showcase.json) are added; --bare writes layouts only.
build_deck.py removes those slides before building a real deck.
"""
import argparse
import re
from pathlib import Path

from pptx import Presentation
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.opc.packuri import PackURI
from pptx.oxml import parse_xml
from pptx.parts.slide import SlideLayoutPart

from rs_style import (PALETTE, FONT_REGULAR, FONT_HEAVY, SLIDE_W, SLIDE_H,
                      emu, clr_xml, fill_xml, line_xml)

ROOT = Path(__file__).resolve().parent.parent
LOGO = ROOT / "assets" / "logos"

NS = ('xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
      'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
      'xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"')

CT_LAYOUT = "application/vnd.openxmlformats-officedocument.presentationml.slideLayout+xml"


# ---------------------------------------------------------------------------
# XML builders
# ---------------------------------------------------------------------------
class Ids:
    def __init__(self):
        self.n = 1

    def __call__(self):
        self.n += 1
        return self.n


def font_xml(heavy=False):
    f = FONT_HEAVY if heavy else FONT_REGULAR
    return f'<a:latin typeface="{f}"/><a:ea typeface="{f}"/><a:cs typeface="{f}"/>'


def ppr(lvl=1, sz=24, color="tx1", heavy=False, bold=False, bullet=None,
        marL=0.0, indent=0.0, algn="l", lnSpc=100, spcBef=0, spcAft=0, tag="lvl"):
    """One <a:lvlNpPr>. bullet: None -> buNone, ('check'|'arrow'|'dash')."""
    if bullet is None:
        bu = "<a:buNone/>"
    else:
        char = {"check": "", "arrow": "", "dash": "–"}[bullet]
        font = "Arial" if bullet == "dash" else "Wingdings"
        bu = (f'<a:buClr>{clr_xml("accent1")}</a:buClr><a:buSzPct val="90000"/>'
              f'<a:buFont typeface="{font}" pitchFamily="2" charset="2"/>'
              f'<a:buChar char="{char}"/>')
    b = ' b="1"' if bold else ""
    return (f'<a:{tag}{lvl}pPr marL="{emu(marL)}" indent="{emu(indent)}" algn="{algn}">'
            f'<a:lnSpc><a:spcPct val="{lnSpc*1000}"/></a:lnSpc>'
            f'<a:spcBef><a:spcPts val="{spcBef*100}"/></a:spcBef>'
            f'<a:spcAft><a:spcPts val="{spcAft*100}"/></a:spcAft>{bu}'
            f'<a:defRPr lang="ko-KR" sz="{sz*100}"{b} kern="0">{fill_xml(color)}{font_xml(heavy)}</a:defRPr>'
            f'</a:{tag}{lvl}pPr>')


def bullet_levels(s1=32, s2=28, s3=24, lnSpc=130, spcBef=6):
    """Lab bullet system: lvl1 check, lvl2 arrow, lvl3 dash."""
    return (ppr(1, s1, "tx1", bullet="check", marL=0.55, indent=-0.55, lnSpc=lnSpc, spcBef=spcBef)
            + ppr(2, s2, "tx1:75", bullet="arrow", marL=1.1, indent=-0.5, lnSpc=lnSpc, spcBef=spcBef)
            + ppr(3, s3, "tx1:75", bullet="dash", marL=1.6, indent=-0.4, lnSpc=lnSpc, spcBef=spcBef))


def xfrm(x, y, w, h):
    return (f'<a:xfrm><a:off x="{emu(x)}" y="{emu(y)}"/>'
            f'<a:ext cx="{emu(w)}" cy="{emu(h)}"/></a:xfrm>')


def body_pr(anchor="t", ins=(0.1, 0.05, 0.1, 0.05), autofit="norm", wrap="square"):
    l, t, r, b = (emu(v) for v in ins)
    fit = {"norm": "<a:normAutofit/>", "none": "<a:noAutofit/>", "shape": "<a:spAutoFit/>"}[autofit]
    return (f'<a:bodyPr wrap="{wrap}" lIns="{l}" tIns="{t}" rIns="{r}" bIns="{b}" '
            f'anchor="{anchor}" rtlCol="0">{fit}</a:bodyPr>')


def ph(ids, name, x, y, w, h, *, type_=None, idx=None, prompt="", lst="",
       fill=None, line=None, anchor="t", ins=(0.1, 0.05, 0.1, 0.05), autofit="norm",
       geom="rect", prompt_lvl=0):
    t = f' type="{type_}"' if type_ else ""
    i = f' idx="{idx}"' if idx is not None else ""
    lvl = f'<a:pPr lvl="{prompt_lvl}"/>' if prompt_lvl else ""
    return (f'<p:sp><p:nvSpPr><p:cNvPr id="{ids()}" name="{name}"/>'
            f'<p:cNvSpPr><a:spLocks noGrp="1"/></p:cNvSpPr>'
            f'<p:nvPr><p:ph{t}{i} hasCustomPrompt="1"/></p:nvPr></p:nvSpPr>'
            f'<p:spPr>{xfrm(x, y, w, h)}<a:prstGeom prst="{geom}"><a:avLst/></a:prstGeom>'
            f'{fill_xml(fill) if fill else ""}{line_xml(line) if line else ""}</p:spPr>'
            f'<p:txBody>{body_pr(anchor, ins, autofit)}<a:lstStyle>{lst}</a:lstStyle>'
            f'<a:p>{lvl}<a:r><a:rPr lang="ko-KR" dirty="0"/><a:t>{prompt}</a:t></a:r></a:p>'
            f'</p:txBody></p:sp>')


def rect(ids, name, x, y, w, h, fill, line=None, geom="rect"):
    return (f'<p:sp><p:nvSpPr><p:cNvPr id="{ids()}" name="{name}"/><p:cNvSpPr/><p:nvPr userDrawn="1"/></p:nvSpPr>'
            f'<p:spPr>{xfrm(x, y, w, h)}<a:prstGeom prst="{geom}"><a:avLst/></a:prstGeom>'
            f'{fill_xml(fill)}{line_xml(line)}</p:spPr>'
            f'<p:txBody><a:bodyPr/><a:lstStyle/><a:p><a:endParaRPr lang="ko-KR"/></a:p></p:txBody></p:sp>')


def hline(ids, name, x, y, w, color="accent1", width_pt=1.0):
    return (f'<p:cxnSp><p:nvCxnSpPr><p:cNvPr id="{ids()}" name="{name}"/><p:cNvCxnSpPr/><p:nvPr userDrawn="1"/></p:nvCxnSpPr>'
            f'<p:spPr>{xfrm(x, y, w, 0)}<a:prstGeom prst="line"><a:avLst/></a:prstGeom>'
            f'{line_xml(color, width_pt)}</p:spPr></p:cxnSp>')


def text(ids, name, x, y, w, h, s, *, sz=20, color="tx1", heavy=False, algn="l",
         anchor="ctr", field=None):
    run = (f'<a:rPr lang="ko-KR" sz="{sz*100}" dirty="0">{fill_xml(color)}{font_xml(heavy)}</a:rPr>')
    if field == "slidenum":
        body = (f'<a:fld id="{{B6F15528-21DE-4FAA-801E-634DDDAF4B2B}}" type="slidenum">'
                f'{run}<a:t>‹#›</a:t></a:fld>')
    else:
        body = f'<a:r>{run}<a:t>{s}</a:t></a:r>'
    return (f'<p:sp><p:nvSpPr><p:cNvPr id="{ids()}" name="{name}"/><p:cNvSpPr txBox="1"/><p:nvPr userDrawn="1"/></p:nvSpPr>'
            f'<p:spPr>{xfrm(x, y, w, h)}<a:prstGeom prst="rect"><a:avLst/></a:prstGeom><a:noFill/></p:spPr>'
            f'<p:txBody>{body_pr(anchor, (0.05, 0, 0.05, 0), "none")}<a:lstStyle/>'
            f'<a:p><a:pPr algn="{algn}"/>{body}</a:p></p:txBody></p:sp>')


def pic(ids, name, rid, x, y, w, h):
    return (f'<p:pic><p:nvPicPr><p:cNvPr id="{ids()}" name="{name}"/>'
            f'<p:cNvPicPr><a:picLocks noChangeAspect="1"/></p:cNvPicPr><p:nvPr userDrawn="1"/></p:nvPicPr>'
            f'<p:blipFill><a:blip r:embed="{rid}"/><a:stretch><a:fillRect/></a:stretch></p:blipFill>'
            f'<p:spPr>{xfrm(x, y, w, h)}<a:prstGeom prst="rect"><a:avLst/></a:prstGeom></p:spPr></p:pic>')


def sp_tree(inner):
    return ('<p:spTree><p:nvGrpSpPr><p:cNvPr id="1" name=""/><p:cNvGrpSpPr/><p:nvPr/></p:nvGrpSpPr>'
            f'<p:grpSpPr>{xfrm(0, 0, 0, 0)}</p:grpSpPr>{inner}</p:spTree>')


# ---------------------------------------------------------------------------
# Geometry shared by content layouts
# ---------------------------------------------------------------------------
G = dict(
    section=(0.63, 0.33, 9.0, 0.45),
    title=(0.63, 0.90, 15.9, 0.95),
    source=(6.4, 10.76, 11.2, 0.50),
    band=(0.25, 7.55, 19.5, 3.05),
    band_top=(0.25, 2.05, 19.5, 2.15),
    zone=(0.75, 2.25, 18.5, 5.10),         # content zone above the band
    zone_full=(0.75, 2.25, 18.5, 8.20),    # content zone without band
    footer_bar=(0, 10.74, SLIDE_W, SLIDE_H - 10.74),
)

BAND_FILL = "accent1:8"          # very light Navy Peony tint
CARD_LINE = "accent1:30"


def common_ph(ids, source=True):
    s = ph(ids, "section", *G["section"], type_="body", idx=10, prompt="Section",
           lst=ppr(1, 22, "tx1:75"), anchor="ctr", ins=(0.05, 0, 0.05, 0))
    s += ph(ids, "title", *G["title"], type_="title", prompt="슬라이드 제목: 한 장의 핵심 메시지",
            anchor="t", ins=(0.05, 0, 0.05, 0))
    if source:
        s += ph(ids, "source", *G["source"], type_="body", idx=11,
                prompt="Author et al., Title, Venue Year", anchor="ctr",
                lst=ppr(1, 16, "bg1", algn="r"), ins=(0.05, 0, 0.1, 0), autofit="none")
    return s


def band_ph(ids, geo="band", prompt_top=False):
    return ph(ids, "band", *G[geo], type_="body", idx=12,
              prompt="핵심 주장을 한 줄로 (명사형 종결)", fill=BAND_FILL, anchor="ctr",
              lst=bullet_levels(32, 28, 24, lnSpc=125, spcBef=4), ins=(0.4, 0.15, 0.4, 0.15))


def fig_ph(ids, name, idx, x, y, w, h):
    return ph(ids, name, x, y, w, h, type_="pic", idx=idx,
              prompt="그림 삽입: 논문 figure·표 캡쳐 (비율 유지, 잘라내지 않음)",
              lst=ppr(1, 20, "accent6", algn="ctr"), anchor="ctr")


# ---------------------------------------------------------------------------
# Layout table. Each entry: (name, showMasterSp, builder(ids, img) -> xml)
# img(name) returns an rId for a logo in that layout part.
# ---------------------------------------------------------------------------
def L_title(ids, img):
    s = ph(ids, "venue_logo", 0.5, 0.45, 2.6, 1.3, type_="pic", idx=13,
           prompt="학회 로고", lst=ppr(1, 14, "accent6", algn="ctr"), anchor="ctr")
    s += hline(ids, "line_l", 3.25, 1.12, 4.7)
    s += hline(ids, "line_r", 12.05, 1.12, 4.75)
    s += ph(ids, "date", 8.05, 0.84, 3.9, 0.56, type_="body", idx=10, prompt="2026년 Q3",
            lst=ppr(1, 22, "tx1:75", algn="ctr"), anchor="ctr", autofit="none")
    s += pic(ids, "logo_snu", img("snu_seal.png"), 17.10, 0.60, 1.07, 1.11)
    s += pic(ids, "logo_dsba", img("dsba_square.png"), 18.31, 0.59, 1.28, 1.12)
    s += ph(ids, "title", 1.5, 2.55, 17.0, 2.45, type_="ctrTitle", prompt="논문 제목",
            lst=ppr(1, 48, "accent1", heavy=True, algn="ctr", lnSpc=110), anchor="b")
    s += ph(ids, "citation", 1.5, 5.20, 17.0, 1.25, type_="subTitle", idx=1,
            prompt="First Author et al., “Paper Title”, Venue, Year, arXiv:XXXX.XXXXX",
            lst=ppr(1, 20, "tx1:75", algn="ctr", lnSpc=120), anchor="t")
    s += ph(ids, "presenter", 4.0, 6.85, 12.0, 2.3, type_="body", idx=11,
            prompt="서울대학교 산업공학과", lst=ppr(1, 24, "tx1", algn="ctr", lnSpc=130),
            anchor="t", autofit="none")
    s += rect(ids, "footer_bar", *G["footer_bar"], "accent1")
    return s


def L_contents(ids, img):
    s = rect(ids, "panel", 0, 0, 6.33, SLIDE_H, "accent1")
    s += text(ids, "contents_label", 0.55, 0.85, 5.5, 1.3, "Contents", sz=66, color="bg1",
              heavy=True, anchor="t")
    s += pic(ids, "logo_dsba_white", img("dsba_white.png"), 0.28, 9.95, 2.42, 1.24)
    tab = (f'<a:tabLst><a:tab pos="{emu(1.4)}" algn="l"/></a:tabLst>')
    lst = ppr(1, 36, "tx1:55", lnSpc=100, spcAft=34).replace("<a:buNone/>", "<a:buNone/>" + tab)
    s += ph(ids, "toc", 7.4, 1.6, 11.6, 8.4, type_="body", idx=12,
            prompt="01\tBackground", lst=lst, anchor="ctr", autofit="none")
    return s


def L_A(ids, img):
    return (common_ph(ids) + fig_ph(ids, "figure", 13, 1.0, 2.25, 18.0, 5.10) + band_ph(ids))


def L_B(ids, img):
    return (common_ph(ids) + fig_ph(ids, "figure_left", 13, 0.75, 2.25, 8.9, 5.10)
            + fig_ph(ids, "figure_right", 14, 10.35, 2.25, 8.9, 5.10) + band_ph(ids))


def L_C(ids, img):
    return (common_ph(ids) + fig_ph(ids, "figure", 13, 0.75, 2.25, 10.4, 5.10)
            + ph(ids, "explain", 11.55, 2.25, 7.7, 5.10, type_="body", idx=15,
                 prompt="그림 해설 (짧은 항목)", lst=bullet_levels(26, 24, 22, lnSpc=120, spcBef=4),
                 anchor="ctr")
            + band_ph(ids))


def L_D(ids, img):
    return (common_ph(ids) + band_ph(ids, "band_top")
            + ph(ids, "wide", 0.75, 4.45, 18.5, 6.0, idx=16,
                 prompt="표·차트·넓은 그림", lst=ppr(1, 20, "accent6", algn="ctr"), anchor="ctr"))


def L_E(ids, img):
    s = common_ph(ids)
    for side, x, i in (("left", 0.75, 17), ("right", 10.25, 19)):
        s += ph(ids, f"header_{side}", x, 2.30, 9.0, 1.0, type_="body", idx=i,
                prompt=f"{'왼쪽' if side == 'left' else '오른쪽'} 제목", fill="accent1",
                lst=ppr(1, 28, "bg1", heavy=True, algn="ctr"), anchor="ctr", autofit="none")
        s += ph(ids, f"body_{side}", x, 3.30, 9.0, 6.95, type_="body", idx=i + 1,
                prompt="항목", line=CARD_LINE, lst=bullet_levels(26, 24, 22, lnSpc=125, spcBef=8),
                anchor="t", ins=(0.4, 0.35, 0.35, 0.3))
    return s


def L_F(ids, img):
    return common_ph(ids) + band_ph(ids)


def L_G(ids, img):
    return common_ph(ids)


def L_H(ids, img):
    return (common_ph(ids)
            + ph(ids, "statement", 1.5, 3.0, 17.0, 3.0, type_="body", idx=21,
                 prompt="핵심 질문 또는 한 문장 요약", anchor="b",
                 lst=ppr(1, 48, "accent1", heavy=True, algn="ctr", lnSpc=115))
            + ph(ids, "support", 2.5, 6.3, 15.0, 3.0, type_="body", idx=22,
                 prompt="보충 설명 1~2줄", anchor="t",
                 lst=ppr(1, 28, "tx1:75", algn="ctr", lnSpc=130)))


def L_closing(ids, img):
    s = hline(ids, "line_l", 0.5, 1.12, 7.45)
    s += hline(ids, "line_r", 12.05, 1.12, 4.75)
    s += ph(ids, "date", 8.05, 0.84, 3.9, 0.56, type_="body", idx=10, prompt="2026년 Q3",
            lst=ppr(1, 22, "tx1:75", algn="ctr"), anchor="ctr", autofit="none")
    s += pic(ids, "logo_snu", img("snu_seal.png"), 17.10, 0.60, 1.07, 1.11)
    s += pic(ids, "logo_dsba", img("dsba_square.png"), 18.31, 0.59, 1.28, 1.12)
    s += ph(ids, "title", 3.5, 3.7, 13.0, 2.3, type_="ctrTitle", prompt="감사합니다",
            lst=ppr(1, 88, "accent1", heavy=True, algn="ctr"), anchor="b", autofit="none")
    s += ph(ids, "subtitle", 3.5, 6.3, 13.0, 0.9, type_="subTitle", idx=1,
            prompt="Papers You Should Read · 2026 Q3",
            lst=ppr(1, 26, "tx1:75", algn="ctr"), anchor="t")
    s += rect(ids, "footer_bar", *G["footer_bar"], "accent1")
    return s


def L_guide(ids, img):
    s = rect(ids, "accent_bar", 0.75, 0.9, 0.08, 0.8, "accent1")
    s += ph(ids, "title", 1.1, 0.85, 14.5, 0.9, type_="title", prompt="가이드 제목",
            lst=ppr(1, 36, "accent1", heavy=True), anchor="ctr", ins=(0.05, 0, 0.05, 0))
    s += text(ids, "tag", 15.8, 1.0, 3.45, 0.6, "숨김 가이드 · 발표에 나오지 않음", sz=16,
              color="accent2", algn="r")
    s += hline(ids, "rule", 0.75, 1.95, 18.5, "accent1:25", 1.0)
    s += ph(ids, "body", 1.1, 2.3, 17.9, 8.2, type_="body", idx=1, prompt="내용",
            lst=bullet_levels(24, 21, 19, lnSpc=120, spcBef=6), anchor="t")
    s += rect(ids, "footer_bar", *G["footer_bar"], "accent1")
    return s


LAYOUTS = [
    # name, show master chrome, builder, one-line purpose (also written to layouts.md)
    ("Title", False, L_title, "표지: 논문 제목, 인용, 발표자, 학회 로고"),
    ("Contents", False, L_contents, "목차 겸 섹션 구분: 현재 섹션만 강조"),
    ("A_FigureFull_Band", True, L_A, "그림 1개 전면 + 하단 밴드"),
    ("B_FigureTwo_Band", True, L_B, "그림 2개 비교 + 하단 밴드"),
    ("C_FigureExplain_Band", True, L_C, "좌 그림 + 우 해설 + 하단 밴드"),
    ("D_TopBand_Wide", True, L_D, "상단 밴드 + 하단 넓은 표·차트"),
    ("E_TwoColumn", True, L_E, "2단 카드 정리 (비교·결론)"),
    ("F_Diagram_Band", True, L_F, "도형·수식 자유 영역 + 하단 밴드"),
    ("G_TitleOnly", True, L_G, "제목만 (자유 구성, 부록)"),
    ("H_Statement", True, L_H, "핵심 질문 또는 한 문장 요약"),
    ("Closing", False, L_closing, "마무리"),
    ("Guide", False, L_guide, "숨김 가이드 슬라이드 전용"),
]


# ---------------------------------------------------------------------------
# Master
# ---------------------------------------------------------------------------
def master_xml(ids, img, layout_rids):
    chrome = rect(ids, "accent_bar", 0.425, 0.26, 0.05, 1.53, "accent1")
    chrome += pic(ids, "logo_kiie", img("kiie_emblem.png"), 16.92, 0.42, 1.36, 1.10)
    chrome += pic(ids, "logo_dsba", img("dsba_square.png"), 18.40, 0.42, 1.26, 1.10)
    chrome += rect(ids, "footer_bar", *G["footer_bar"], "accent1")
    chrome += text(ids, "footer_text", 0.1, 10.76, 6.3, 0.5, "Papers You Should Read  ·  DSBA Lab",
                   sz=20, color="bg1")
    chrome += text(ids, "slide_number", 17.7, 10.76, 2.1, 0.5, "", sz=22, color="bg1",
                   algn="r", field="slidenum")
    phs = ph(ids, "title", *G["title"], type_="title", prompt="마스터 제목 스타일",
             anchor="t", ins=(0.05, 0, 0.05, 0))
    phs += ph(ids, "body", *G["zone_full"], type_="body", idx=1, prompt="마스터 본문 스타일",
              anchor="t")
    lids = "".join(f'<p:sldLayoutId id="{2147483649+i}" r:id="{rid}"/>'
                   for i, rid in enumerate(layout_rids))
    title_style = ppr(1, 44, "accent1", heavy=True, lnSpc=100)
    body_style = bullet_levels(32, 28, 24) + "".join(
        ppr(n, 20, "tx1:75", marL=1.6 + 0.4 * (n - 3), indent=0) for n in range(4, 10))
    other = ppr(1, 18, "tx1")
    return (f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            f'<p:sldMaster {NS}><p:cSld><p:bg><p:bgPr>{fill_xml("bg1")}<a:effectLst/></p:bgPr></p:bg>'
            f'{sp_tree(chrome + phs)}</p:cSld>'
            f'<p:clrMap bg1="lt1" tx1="dk1" bg2="lt2" tx2="dk2" accent1="accent1" accent2="accent2" '
            f'accent3="accent3" accent4="accent4" accent5="accent5" accent6="accent6" hlink="hlink" folHlink="folHlink"/>'
            f'<p:sldLayoutIdLst>{lids}</p:sldLayoutIdLst>'
            f'<p:hf hdr="0" ftr="0" dt="0"/>'
            f'<p:txStyles><p:titleStyle>{title_style}</p:titleStyle>'
            f'<p:bodyStyle>{body_style}</p:bodyStyle>'
            f'<p:otherStyle>{other}</p:otherStyle></p:txStyles></p:sldMaster>')


def layout_xml(name, show_master, inner):
    sm = "" if show_master else ' showMasterSp="0"'
    return (f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            f'<p:sldLayout {NS} preserve="1" userDrawn="1"{sm}><p:cSld name="{name}">'
            f'{sp_tree(inner)}</p:cSld><p:clrMapOvr><a:masterClrMapping/></p:clrMapOvr></p:sldLayout>')


def theme_xml(base: bytes) -> bytes:
    s = base.decode("utf-8")
    order = ["dk1", "lt1", "dk2", "lt2", "accent1", "accent2", "accent3", "accent4",
             "accent5", "accent6", "hlink", "folHlink"]
    clr = "".join(f'<a:{k}><a:srgbClr val="{PALETTE[k][0]}"/></a:{k}>' for k in order)
    s = re.sub(r'<a:clrScheme name="[^"]*">.*?</a:clrScheme>',
               f'<a:clrScheme name="research-slides Pantone A">{clr}</a:clrScheme>', s, flags=re.S)
    fnt = lambda f: (f'<a:latin typeface="{f}"/><a:ea typeface="{f}"/><a:cs typeface=""/>'
                     f'<a:font script="Hang" typeface="{f}"/>')
    s = re.sub(r'<a:fontScheme name="[^"]*">.*?</a:fontScheme>',
               f'<a:fontScheme name="research-slides"><a:majorFont>{fnt(FONT_HEAVY)}</a:majorFont>'
               f'<a:minorFont>{fnt(FONT_REGULAR)}</a:minorFont></a:fontScheme>', s, flags=re.S)
    s = re.sub(r'<a:theme ([^>]*)name="[^"]*"', r'<a:theme \1name="research-slides"', s)
    return s.encode("utf-8")


# ---------------------------------------------------------------------------
def build(out: Path):
    prs = Presentation()
    prs.slide_width, prs.slide_height = emu(SLIDE_W), emu(SLIDE_H)
    master = prs.slide_masters[0]
    mpart = master.part
    pkg = mpart.package

    # theme
    theme = mpart.part_related_by(RT.THEME)
    theme._blob = theme_xml(theme.blob)

    # drop default layouts
    for lid in list(master._element.sldLayoutIdLst):
        mpart.drop_rel(lid.rId)

    # new layouts
    rids = []
    for n, (name, show, fn, _) in enumerate(LAYOUTS, start=1):
        el = parse_xml(layout_xml(name, show, "").encode())
        part = SlideLayoutPart(PackURI(f"/ppt/slideLayouts/slideLayout{n}.xml"), CT_LAYOUT, pkg, el)
        part.relate_to(mpart, RT.SLIDE_MASTER)
        rids.append(mpart.relate_to(part, RT.SLIDE_LAYOUT))  # relate first so image names are unique
        cache = {}

        def img(fname, part=part, cache=cache):
            if fname not in cache:
                cache[fname] = part.get_or_add_image_part(str(LOGO / fname))[1]
            return cache[fname]

        inner = fn(Ids(), img)
        part._element = parse_xml(layout_xml(name, show, inner).encode())

    mcache = {}

    def mimg(fname):
        if fname not in mcache:
            mcache[fname] = mpart.get_or_add_image_part(str(LOGO / fname))[1]
        return mcache[fname]

    mpart._element = parse_xml(master_xml(Ids(), mimg, rids).encode())

    cp = prs.core_properties
    cp.title = "research-slides template (DSBA Papers You Should Read)"
    cp.subject = "Build decks from this file with the research-slides skill."
    cp.keywords = "research-slides; DSBA; PYSR; claude-skill"
    cp.comments = ("For Claude: use the research-slides skill (SKILL.md) and "
                   "scripts/build_deck.py. Do not hand-place shapes.")
    cp.author = "DSBA Lab"
    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(out)
    # reopen to make sure the package is valid
    Presentation(out)
    print(f"wrote {out} with {len(LAYOUTS)} layouts")


if __name__ == "__main__":
    import tempfile
    import build_deck
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", default=str(ROOT / "assets" / "research-slides-template.pptx"))
    ap.add_argument("--bare", action="store_true", help="layouts only, no guide/example slides")
    a = ap.parse_args()
    if a.bare:
        build(Path(a.out))
    else:
        with tempfile.TemporaryDirectory() as d:
            bare = Path(d) / "bare.pptx"
            build(bare)
            build_deck.build(ROOT / "examples" / "template_showcase.json", a.out, template=bare)
