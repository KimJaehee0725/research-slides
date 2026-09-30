"""Native PowerPoint equations (OMML) from LaTeX.

PowerPoint stores equations as Office Math (OMML) inside `a14:m`, wrapped in
`mc:AlternateContent` so that apps without equation support show a fallback.
We convert LaTeX -> OMML with pandoc (the same converter pandoc uses for
.docx), then restyle the runs the way PowerPoint expects (a:rPr + Cambria Math).

Two entry points:
  * equation_shape(slide, x, y, w, h, tex, size)  -> a standalone equation box
      fallback = PNG rendered with matplotlib (so non-Office viewers still show it)
  * inline math inside any text: write `$...$` in spec strings; components.add_runs
    calls omath_element() and build_deck wraps the shape with a plain-text fallback.
"""
import copy
import functools
import re
import shutil
import subprocess
import tempfile
import zipfile
from pathlib import Path

from lxml import etree

NS_M = "http://schemas.openxmlformats.org/officeDocument/2006/math"
NS_A = "http://schemas.openxmlformats.org/drawingml/2006/main"
NS_A14 = "http://schemas.microsoft.com/office/drawing/2010/main"
NS_MC = "http://schemas.openxmlformats.org/markup-compatibility/2006"
NS_P = "http://schemas.openxmlformats.org/presentationml/2006/main"
MATH_FONT = "Cambria Math"


def _pandoc():
    exe = shutil.which("pandoc")
    if exe:
        return exe
    try:
        import pypandoc
        return pypandoc.get_pandoc_path()
    except Exception:
        raise SystemExit("pandoc is required for equations: install pandoc or `pip install pypandoc-binary`")


@functools.lru_cache(maxsize=512)
def _omml_xml(tex: str, display: bool) -> bytes:
    """LaTeX -> <m:oMathPara> (display) or <m:oMath> (inline) XML bytes."""
    with tempfile.TemporaryDirectory() as d:
        md = Path(d) / "m.md"
        md.write_text(f"$${tex}$$\n" if display else f"x ${tex}$ x\n", encoding="utf-8")
        out = Path(d) / "m.docx"
        subprocess.run([_pandoc(), str(md), "-o", str(out)], check=True,
                       stdout=subprocess.DEVNULL, stderr=subprocess.PIPE)
        doc = zipfile.ZipFile(out).read("word/document.xml")
    root = etree.fromstring(doc)
    tag = "oMathPara" if display else "oMath"
    el = root.find(f".//{{{NS_M}}}{tag}")
    if el is None:
        raise SystemExit(f"pandoc could not parse the formula: {tex}")
    return etree.tostring(el)


def _style_runs(el, size_pt=None, color_xml=None, bold=False):
    """Give every math run a DrawingML rPr (PowerPoint needs a:rPr, not w:rPr)."""
    for r in el.iter(f"{{{NS_M}}}r"):
        mrpr = r.find(f"{{{NS_M}}}rPr")
        plain = mrpr is not None and mrpr.find(f"{{{NS_M}}}sty") is not None \
            and mrpr.find(f"{{{NS_M}}}sty").get(f"{{{NS_M}}}val") == "p"
        rpr = etree.Element(f"{{{NS_A}}}rPr", lang="en-US", i="0" if plain else "1",
                            b="1" if bold else "0")
        if size_pt:
            rpr.set("sz", str(int(size_pt * 100)))
        if color_xml is not None:
            sf = etree.SubElement(rpr, f"{{{NS_A}}}solidFill")
            sf.append(copy.deepcopy(color_xml))
        etree.SubElement(rpr, f"{{{NS_A}}}latin", typeface=MATH_FONT)
        etree.SubElement(rpr, f"{{{NS_A}}}cs", typeface=MATH_FONT)
        idx = 1 if mrpr is not None else 0
        r.insert(idx, rpr)
    return el


def omath_element(tex, display=False, size_pt=None, color_xml=None, align="centerGroup"):
    """Return an <a14:m> element holding the equation (ready to put in <a:p>)."""
    el = etree.fromstring(_omml_xml(tex, display))
    if display:
        jc = el.find(f".//{{{NS_M}}}jc")
        if jc is not None:
            jc.set(f"{{{NS_M}}}val", align)
    _style_runs(el, size_pt, color_xml)
    etree.cleanup_namespaces(el, top_nsmap={"m": NS_M, "a": NS_A})
    m = etree.Element(f"{{{NS_A14}}}m", nsmap={"a14": NS_A14, "m": NS_M})
    m.append(el)
    return m


def linear_text(tex):
    """Readable plain-text version of a LaTeX formula (for text fallbacks)."""
    s = tex
    rep = {r"\ell": "ℓ", r"\sigma": "σ", r"\alpha": "α", r"\beta": "β", r"\theta": "θ",
           r"\pi": "π", r"\phi": "φ", r"\lambda": "λ", r"\mu": "μ", r"\dagger": "†",
           r"\partial": "∂", r"\ge": "≥", r"\le": "≤", r"\neq": "≠", r"\approx": "≈",
           r"\cdot": "·", r"\times": "×", r"\to": "→", r"\rightarrow": "→", r"\infty": "∞",
           r"\sum": "Σ", r"\prime": "′", r"\quad": "  ", r"\,": " ", r"\left": "",
           r"\right": "", r"\mathbb{E}": "E", r"\leftarrow": "←", r"\Rightarrow": "⇒", r"\leq": "≤", r"\geq": "≥"}
    for k in sorted(rep, key=len, reverse=True):  # \left before \le, etc.
        s = s.replace(k, rep[k])
    s = re.sub(r"\\(?:text|mathrm|mathbf|operatorname)\{([^}]*)\}", r"\1", s)
    s = re.sub(r"\\frac\{([^}]*)\}\{([^}]*)\}", r"(\1)/(\2)", s)
    s = s.replace("{", "").replace("}", "").replace("\\", "")
    return s


def _mpl_png(tex, size_pt, color="#262626"):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    t = re.sub(r"\\text\{", r"\\mathrm{", tex).replace(r"\quad", r"\ \ ")
    for a, b in ((r"\geq", r"\geq"), (r"\ge ", r"\geq "), (r"\le ", r"\leq "),
                 (r"\operatorname{", r"\mathrm{")):
        t = t.replace(a, b)
    t = re.sub(r"\\ge(?![a-z])", r"\\geq", t)
    t = re.sub(r"\\le(?![a-z])", r"\\leq", t)
    fig = plt.figure(figsize=(0.01, 0.01))
    fig.text(0, 0, f"${t}$", fontsize=size_pt, color=color)
    out = Path(tempfile.mkstemp(suffix=".png")[1])
    fig.savefig(out, dpi=300, transparent=True, bbox_inches="tight", pad_inches=0.04)
    plt.close(fig)
    return out


def _alternate(choice_el, fallback_el):
    ac = etree.Element(f"{{{NS_MC}}}AlternateContent", nsmap={"mc": NS_MC})
    ch = etree.SubElement(ac, f"{{{NS_MC}}}Choice", nsmap={"a14": NS_A14}, Requires="a14")
    ch.append(choice_el)
    fb = etree.SubElement(ac, f"{{{NS_MC}}}Fallback")
    fb.append(fallback_el)
    return ac


def equation_shape(slide, x, y, w, h, tex, size=32, align="c", color_xml=None):
    """Standalone equation text box with an image fallback. Returns the AlternateContent."""
    from pptx.util import Inches
    from pptx.enum.text import MSO_ANCHOR
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    jc = {"c": "centerGroup", "l": "left", "r": "right"}[align]
    p._p.insert(0, omath_element(tex, display=True, size_pt=size, color_xml=color_xml, align=jc))
    sp = tb._element

    # fallback picture of the same formula, fitted in the same box
    fallback = None
    try:
        png = _mpl_png(tex, size)
        from PIL import Image
        with Image.open(png) as im:
            iw, ih = im.size
        scale = min(w / (iw / 300), h / (ih / 300), 1.0)
        dw, dh = iw / 300 * scale, ih / 300 * scale
        dx = x + {"l": 0, "c": (w - dw) / 2, "r": w - dw}[align]
        pic = slide.shapes.add_picture(str(png), Inches(dx), Inches(y + (h - dh) / 2),
                                       Inches(dw), Inches(dh))
        fallback = pic._element
        fallback.getparent().remove(fallback)
    except Exception:
        fallback = None
    if fallback is None:  # plain-text fallback
        fallback = copy.deepcopy(sp)
        for m in fallback.iter(f"{{{NS_A14}}}m"):
            _replace_with_text(m, tex)

    parent = sp.getparent()
    i = parent.index(sp)
    parent.remove(sp)
    ac = _alternate(sp, fallback)
    parent.insert(i, ac)
    return ac


def _replace_with_text(m_el, tex=None):
    txt = tex if tex is not None else m_el.get("data-tex", "")
    r = etree.Element(f"{{{NS_A}}}r")
    etree.SubElement(r, f"{{{NS_A}}}rPr", lang="en-US")
    t = etree.SubElement(r, f"{{{NS_A}}}t")
    t.text = linear_text(txt)
    m_el.getparent().replace(m_el, r)


def wrap_inline_math(slide):
    """Wrap every shape that contains inline <a14:m> in AlternateContent,
    with a fallback copy where each equation is plain text."""
    tree = slide.shapes._spTree
    for sp in list(tree):
        if sp.tag != f"{{{NS_P}}}sp":
            continue
        ms = list(sp.iter(f"{{{NS_A14}}}m"))
        if not ms:
            continue
        fb = copy.deepcopy(sp)
        for m in list(fb.iter(f"{{{NS_A14}}}m")):
            _replace_with_text(m)
        for m in ms:  # the tex was only needed for the fallback
            m.attrib.pop("data-tex", None)
        i = tree.index(sp)
        tree.remove(sp)
        tree.insert(i, _alternate(sp, fb))
